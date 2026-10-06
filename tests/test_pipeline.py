"""End-to-end run on the fake web: balances, workbook formulas, and the source-off rule."""
import shutil
import subprocess
from datetime import date

import pandas as pd
import pytest
from openpyxl import load_workbook

from latam_gas_snd import config
from latam_gas_snd.model import MEMO
from latam_gas_snd.pipeline import run
from tests import fakes

START, END = date(2021, 1, 1), date(2021, 1, 14)


def _line(bals, country, key):
    return next(ln for ln in bals[country].lines if ln.key == key)


@pytest.fixture
def full_run(monkeypatch, tmp_path):
    fakes.install(monkeypatch)
    paths = run(START, END, off=(), out_dir=tmp_path / "out")
    return paths, run.last_context, run.last_balances


def test_all_sources_fetched(full_run):
    _, ctx, _ = full_run
    status = {s: r.status for s, r in ctx.results.items()}
    assert status.pop("kpler") == "no_credentials"
    assert set(status.values()) == {"ok"}, status


def test_argentina_lines(full_run):
    _, ctx, bals = full_run
    i = 5
    day = pd.Timestamp(START) + pd.Timedelta(days=i)
    # chart in MMm3/d: pipelines 20 + 30, GNL Escobar series excluded from production
    assert _line(bals, "Argentina", "ar_prod_inj").series[day] == pytest.approx(50.0, abs=1e-3)
    esc = fakes.partes_value("importaciones", "GNL Escobar", i) / 1000
    assert _line(bals, "Argentina", "ar_lng_escobar").series[day] == pytest.approx(esc)
    meth = fakes.partes_value("exp_fuera", "Methanex PAN", i) / 1000
    assert _line(bals, "Chile", "cl_methanex").series[day] == pytest.approx(meth)
    prio = next(ln for ln in bals["Argentina"].lines if ln.label == "Demand: Priority (residential, commercial)")
    assert prio.series[day] == pytest.approx(25.0, abs=1e-3)
    lp = _line(bals, "Argentina", "ar_linepack").series
    assert lp[pd.Timestamp(START)] != lp[pd.Timestamp(START)]           # first day: no previous level -> blank
    assert lp[day] == pytest.approx(fakes.LP_LEVEL[i] - fakes.LP_LEVEL[i - 1], abs=1e-3)
    assert not [c for c in ctx.checks if c.severity == "error"]
    # acceptance: chart series sum to the chart total (all series, incl. the excluded LNG one)
    for sid in ("enargas_injection", "enargas_demand"):
        summary = next(c for c in ctx.checks if c.check == f"{sid}: chart series sum vs chart total")
        assert summary.severity == "info" and summary.value == 0
    assert _line(bals, "Argentina", "ar_memo_chart_total_injection").series[day] == pytest.approx(55.0, abs=1e-3)
    assert not [c for c in ctx.checks if "matched no marks" in c.detail and "Mínimo" not in c.detail]


def test_brazil_lines_and_bolivia_split(full_run):
    _, ctx, bals = full_run
    day = pd.Timestamp("2021-01-07")
    i = (day - pd.Timestamp(START)).days
    assert _line(bals, "Brazil", "br_production").series[day] == pytest.approx(40.0)
    assert _line(bals, "Brazil", "br_lng").series[day] == pytest.approx(10.0)
    via = fakes.partes_value("exp_dentro", "Por Bolivia (TGN)", i) / 1000
    assert _line(bals, "Brazil", "br_imp_ar_via_bo").series[day] == pytest.approx(via)
    assert _line(bals, "Brazil", "br_imp_bolivia").series[day] == pytest.approx(20.0 - via)
    assert _line(bals, "Bolivia", "bo_exp_br").series[day] == pytest.approx(20.0 - via)
    assert _line(bals, "Brazil", "br_memo_transfers").series[day] == pytest.approx(0.0)
    assert _line(bals, "Brazil", "br_unc_delivery").series[day] == pytest.approx(0.7)
    assert _line(bals, "Brazil", "br_linepack").series[day] == pytest.approx(-0.1)
    # ONS: 500 MW x 24h -> gas equivalent
    assert _line(bals, "Brazil", "br_memo_ons").series[day] == pytest.approx(config.ons_mwh_to_mcm(12000))
    checks = [c.check for c in ctx.checks]
    assert any("not classified" in c for c in checks)
    assert any("unit inferred" in c for c in checks)


def test_colombia_lines(full_run):
    _, _, bals = full_run
    day = pd.Timestamp("2021-01-03")
    f = config.MBTU_TO_MCM
    assert _line(bals, "Colombia", "co_prod").series[day] == pytest.approx(900000 * f)
    assert _line(bals, "Colombia", "co_lng").series[day] == pytest.approx(200000 * f)
    power = next(ln for ln in bals["Colombia"].lines if ln.label == "Demand: Power")
    assert power.series[day] == pytest.approx(250000 * f)
    xm_line = _line(bals, "Colombia", "co_xm")
    assert xm_line.side == MEMO                     # BMC has a thermal split
    assert xm_line.series[day] == pytest.approx(100000 * f)   # gas only, ACPM excluded


def test_outputs_exist(full_run):
    paths, _, _ = full_run
    for k in ("main", "Argentina", "Brazil", "Colombia", "brazil_points"):
        assert paths[k].exists()
    wb = load_workbook(paths["main"], read_only=True)
    assert {"Argentina", "Brazil", "Colombia", "Bolivia", "Chile", "Uruguay", "Flows by route", "Sources",
            "Coverage", "Run log", "Checks"} <= set(wb.sheetnames)
    pts = pd.read_parquet(paths["brazil_points"])
    assert {"date", "point", "category", "value"} <= set(pts.columns)
    br = load_workbook(paths["Brazil"], read_only=True)
    assert {"Balance", "By category", "By category & state", "By point", "Point list", "BR checks"} <= set(br.sheetnames)
    co = load_workbook(paths["Colombia"], read_only=True)
    assert {"Balance", "Injection by entry point", "Offtake by sector", "XM burn by plant"} <= set(co.sheetnames)


def test_formulas_and_provenance_rows(full_run):
    paths, _, _ = full_run
    ws = load_workbook(paths["main"])["Uruguay"]
    header = [c.value for c in ws[2]]
    assert header[0] == "Date"
    ts = header.index("Total supply") + 1
    r = 7
    assert str(ws.cell(r, ts).value).startswith("=SUM(")
    assert str(ws.cell(3, 2).value).startswith("ENARGAS partes")          # source under the column
    assert ws.cell(4, 2).value == "daily"                                  # frequency under the column
    assert str(ws.cell(r, header.index("Missing inputs (blank supply/demand cells)") + 1).value).startswith("=COUNTBLANK(")


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_formulas_evaluate(full_run, tmp_path):
    paths, _, bals = full_run
    out = tmp_path / "recalc"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out),
                    str(paths["main"])], check=True, capture_output=True, timeout=180)
    wb = load_workbook(out / paths["main"].name, data_only=True)
    for country in ("Argentina", "Brazil", "Chile", "Uruguay", "Bolivia"):
        ws = wb[country]
        header = [c.value for c in ws[2]]
        res_col = next(i for i, h in enumerate(header) if h and str(h).startswith("Residual")) + 1
        miss_col = header.index("Missing inputs (blank supply/demand cells)") + 1
        bal = bals[country]
        expected = bal.residual()
        for i in range(len(bal.index)):
            got = ws.cell(7 + i, res_col).value
            assert got == pytest.approx(expected.iloc[i], abs=1e-6), (country, i)
        n_inputs = len([ln for ln in bal.lines if ln.side != MEMO])
        blanks = sum(1 for ln in bal.lines if ln.side != MEMO and not ln.formula and pd.isna(ln.series.iloc[0]))
        derived_blank = sum(1 for ln in bal.lines if ln.formula and
                            all(pd.isna(x.series.iloc[0]) for x in bal.lines if x.side == "supply"))
        assert ws.cell(7, miss_col).value == blanks + derived_blank, country
        assert n_inputs >= blanks
    # Bolivia: production and domestic demand have no public daily source -> always missing
    ws = wb["Bolivia"]
    header = [c.value for c in ws[2]]
    assert ws.cell(10, header.index("Missing inputs (blank supply/demand cells)") + 1).value >= 2


def test_source_off_blanks_only_its_lines(monkeypatch, tmp_path):
    """Acceptance: a source switched off leaves its lines blank, is in the run log, changes nothing else."""
    fakes.install(monkeypatch)
    run(START, END, off=(), out_dir=tmp_path / "a")
    base = {(c, ln.key): ln for c, b in run.last_balances.items() for ln in b.lines}
    for off in ("enargas_exp_dentro", "anp", "bmc_offtake_snt", "xm"):
        run(START, END, off=(off,), out_dir=tmp_path / off)
        ctx, bals = run.last_context, run.last_balances
        assert ctx.results[off].status == "off"
        assert any(e.step == off and e.status == "off" for e in ctx.log)
        for c, b in bals.items():
            for ln in b.lines:
                ref = base.get((c, ln.key))
                if off in ln.sources:
                    assert ln.series.isna().all(), (off, c, ln.label)
                elif ref is not None and not ln.formula:
                    pd.testing.assert_series_equal(ln.series, ref.series, check_names=False, obj=f"{off} {c} {ln.label}")
                    assert ln.side == ref.side, (off, ln.label)
